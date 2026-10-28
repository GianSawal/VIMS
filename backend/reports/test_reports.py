"""Phase 13: every report runs and exports, totals match the list summaries, office scope and permissions hold."""
import io

import pytest
from openpyxl import load_workbook

from audit.models import AuditLog
from dashboard.test_dashboard import fleet, fo1, fo2, setup, today  # noqa: F401 (shared fixtures)
from fleet.test_vehicles import as_role
from reports.catalog import REPORTS

pytestmark = pytest.mark.django_db


def run(c, slug, **params):
    r = c.get(f'/api/reports/{slug}/', params)
    assert r.status_code == 200, (slug, r.content[:300])
    return r.json() if not params.get('export') else r


def total(d, key):
    return d['totals'][key]


def test_catalog_follows_permissions(fleet):
    admin = as_role('System Administrator').get('/api/reports/').json()
    assert {r['slug'] for r in admin} == set(REPORTS) and len(REPORTS) == 16
    lto = next(r for r in admin if r['slug'] == 'lto-expiration')
    assert lto['statuses'][0] == {'value': 'EXPIRED', 'label': 'Expired'} and 'period' in lto['filters']
    from accounts.models import User
    from django.contrib.auth.models import Permission
    from rest_framework.test import APIClient
    u = User.objects.create(username='limited')
    u.user_permissions.add(Permission.objects.get(codename='view_vehicle'))
    c = APIClient()
    c.force_authenticate(u)
    assert {r['slug'] for r in c.get('/api/reports/').json()} == {'vehicle-master-list', 'vehicle-status', 'vehicle-history'}
    assert c.get('/api/reports/fuel-cost/').status_code == 403


@pytest.mark.parametrize('slug', sorted(REPORTS))
def test_every_report_runs_and_exports(fleet, slug):
    c = as_role('System Administrator')
    params = {'vehicle': fleet['a'].pk} if slug == 'vehicle-history' else {}
    d = run(c, slug, **params)
    assert d['columns'] and d['count'] == len(d['rows'])
    assert set(d['meta']) >= {'filters', 'statuses', 'group_by'}
    r = run(c, slug, export='csv', **params)
    assert r['Content-Type'].startswith('text/csv')
    lines = r.content.decode('utf-8-sig').strip().splitlines()
    assert lines[0].split(',')[0] == d['columns'][0]['label'].split(',')[0]
    assert len(lines) == 1 + d['count'] + (1 if d['totals'] else 0)
    r = run(c, slug, export='xlsx', **params)
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws['A2'].value == d['title']
    r = run(c, slug, export='pdf', **params)
    assert r['Content-Type'] == 'application/pdf' and r.content.startswith(b'%PDF')
    assert AuditLog.objects.filter(action='export', entity_type='report').count() == 3


def test_totals_match_list_summaries(fleet):
    c = as_role('System Administrator')
    first, last = today().replace(day=1).isoformat(), today().isoformat()
    period = {'date_from': first, 'date_to': last}
    fuel = c.get('/api/fuel-records/summary/', period).json()
    assert total(run(c, 'fuel-cost', **period), 'amount') == fuel['amount'] == '4000.00'
    assert total(run(c, 'fuel-consumption', **period), 'amount') == fuel['amount']
    maint = c.get('/api/maintenance/summary/', period).json()
    assert total(run(c, 'maintenance-expense', **period), 'total_cost') == maint['cost'] == '1500.00'
    assert total(run(c, 'maintenance-history', **period), 'total_cost') == maint['cost']
    by_cat = run(c, 'maintenance-expense', group_by='category', **period)
    assert [r['group'] for r in by_cat['rows']] == ['Engine'] and by_cat['totals']['jobs'] == 1
    trips = c.get('/api/trips/', {**period, 'page_size': 100}).json()
    assert run(c, 'trip-history', **period)['count'] == trips['count'] == 3
    assert total(run(c, 'trip-history', **period), 'distance_km') == 200
    vehicles = c.get('/api/vehicles/summary/').json()
    assert run(c, 'vehicle-master-list')['count'] == sum(vehicles['by_status'].values()) == 4
    util = run(c, 'vehicle-utilization', **period)
    assert (util['totals']['trips'], util['totals']['distance_km']) == (3, 200)
    assert util['rows'][0]['plate_number'] == 'AAA 1' and util['rows'][0]['days_used'] == 1
    office = run(c, 'office-fleet-summary', **period)
    assert office['totals']['total'] == 4 and office['totals']['fuel_cost'] == '4000.00'
    assert office['totals']['maintenance_cost'] == '1500.00' and office['totals']['trips'] == 3


def test_document_and_license_reports(fleet):
    c = as_role('System Administrator')
    lto = run(c, 'lto-expiration')
    assert [r['status'] for r in lto['rows']][:2] == ['Expired', 'Due within 15 days']
    assert run(c, 'lto-expiration', status='EXPIRED')['count'] == 1
    assert run(c, 'insurance-expiration')['rows'][0]['status'] == 'Due within 60 days'
    lic = run(c, 'driver-license-expiration')
    assert {s['label']: s['value'] for s in lic['summary']} == {'Expired': 1, 'Expiring': 1, 'Valid': 0}
    assert run(c, 'driver-license-expiration', status='EXPIRING')['rows'][0]['full_name'] == 'Pedro'


def test_vehicle_history_and_fuel_po(fleet):
    c = as_role('System Administrator')
    assert c.get('/api/reports/vehicle-history/').status_code == 400  # vehicle required
    h = run(c, 'vehicle-history', vehicle=fleet['c'].pk)
    kinds = {r['kind'] for r in h['rows']}
    assert {'Document', 'Fuel', 'Maintenance', 'Trip'} <= kinds
    assert h['totals']['amount'] == '4000.00'  # 2500 fuel + 1500 completed maintenance
    fleet['a'].fuel_records.update(po_number='PO-1')
    po = run(c, 'fuel-po')
    assert po['count'] == 1 and po['rows'][0]['po_number'] == 'PO-1' and po['totals']['amount'] == '1500.00'


def test_field_user_scope(fleet, fo1, fo2):
    c = as_role('Field Office User', [fo1])
    assert {r['office'] for r in run(c, 'vehicle-master-list')['rows']} == {'FO1'}
    assert total(run(c, 'fuel-cost'), 'amount') == '1500.00'
    assert c.get('/api/reports/vehicle-history/', {'vehicle': fleet['c'].pk}).status_code == 404
    assert c.get('/api/reports/trip-history/', {'driver': fleet['d2'].pk}).status_code == 404
    assert run(c, 'vehicle-master-list', office=fo2.pk)['count'] == 0  # asking for another office widens nothing
    assert [r['code'] for r in run(c, 'office-fleet-summary')['rows']] == ['FO1']


def test_filters_and_validation(fleet, fo2):
    c = as_role('System Administrator')
    d = run(c, 'fuel-cost', office=fo2.pk, group_by='vehicle')
    assert [r['group'] for r in d['rows']] == ['BBB 1']
    assert {f['label'] for f in d['filters']} == {'Office', 'Grouped by', 'Date'}
    assert d['params']['date_from'] == today().replace(day=1).isoformat()  # default period: this month
    assert run(c, 'trip-history', status='COMPLETED', vehicle=fleet['a'].pk)['count'] == 2
    assert run(c, 'trip-history', driver=fleet['d2'].pk)['count'] == 1
    one = run(c, 'driver-assignment', driver=fleet['d2'].pk)  # was a 500: filtering the driver list by pk
    assert one['count'] == 1 and {f['label'] for f in one['filters']} >= {'Driver'}
    for bad in ({'status': 'NOPE'}, {'date_from': '2026-05-02', 'date_to': '2026-05-01'}, {'office': 'x'},
                {'group_by': 'driver'}, {'export': 'doc'}):
        assert c.get('/api/reports/fuel-cost/' if 'group_by' in bad else '/api/reports/trip-history/', bad).status_code == 400, bad
    assert c.get('/api/reports/nope/').status_code == 404


def test_logged_out_gets_401(api):
    assert api.get('/api/reports/').status_code == 401
    assert api.get('/api/reports/vehicle-master-list/').status_code == 401
