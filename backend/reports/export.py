"""
Render a built report (see views.build_document) as CSV, Excel or PDF. All three show the same rows and totals;
Excel keeps numbers and dates as real cells so they can be summed, CSV is plain data (UTF-8 with BOM for Excel).
"""
import csv
import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.http import HttpResponse
from django.utils import timezone

BRAND = '0305BA'
LOGO = Path(__file__).parent / 'assets' / 'dole-logo.png'
AGENCY = ('Republic of the Philippines', 'Department of Labor and Employment')
NUMERIC = {'int', 'km', 'money', 'qty', 'num', 'pct'}


def text(value, kind, pretty=False):
    """One cell as text. pretty=True (PDF) uses thousands separators and readable dates."""
    if value is None or value == '':
        return ''
    if kind == 'money':
        return f'{Decimal(value):,.2f}' if pretty else f'{Decimal(value):.2f}'
    if kind == 'qty':
        return f'{Decimal(value):,.3f}' if pretty else f'{Decimal(value):.3f}'
    if kind == 'num':
        return f'{Decimal(value):,.2f}' if pretty else f'{Decimal(value):.2f}'
    if kind == 'pct':
        return f'{Decimal(value):.1f}%'
    if kind in ('int', 'km'):
        return f'{int(value):,}' if pretty else str(int(value))
    if kind == 'datetime' and isinstance(value, datetime):
        return value.strftime('%b %d, %Y %I:%M %p') if pretty else value.strftime('%Y-%m-%d %H:%M')
    if kind == 'date' and isinstance(value, date):
        return value.strftime('%b %d, %Y') if pretty else value.isoformat()
    return str(value)


def filename(doc, ext):
    return f'{doc["slug"]}-{timezone.localdate():%Y%m%d}.{ext}'


def _download(content, content_type, name):
    response = HttpResponse(content, content_type=content_type)
    response['Content-Disposition'] = f'attachment; filename="{name}"'
    return response


def _total_row(doc):
    """Cells of the totals row, 'TOTAL' in the first column; None when the report has no totals."""
    if not doc['totals']:
        return None
    first = doc['columns'][0]['key']
    return [doc['totals'].get(c['key']) if c['key'] != first else 'TOTAL' for c in doc['columns']]


def _meta_lines(doc):
    lines = [f'{label}: {value}' for label, value in doc['filters']]
    if doc['summary']:
        lines.append(' · '.join(f'{s["label"]}: {s["value"]}' for s in doc['summary']))
    return lines


def to_csv(doc):
    out = io.StringIO()
    w = csv.writer(out)
    kinds = [c['kind'] for c in doc['columns']]
    w.writerow([c['label'] for c in doc['columns']])
    for row in doc['rows']:
        w.writerow([text(row.get(c['key']), c['kind']) for c in doc['columns']])
    total = _total_row(doc)
    if total:
        w.writerow([v if i == 0 else text(v, kinds[i]) for i, v in enumerate(total)])
    return _download('﻿' + out.getvalue(), 'text/csv; charset=utf-8', filename(doc, 'csv'))


EXCEL_FORMATS = {'money': '#,##0.00', 'qty': '#,##0.000', 'num': '#,##0.00', 'int': '#,##0', 'km': '#,##0',
                 'pct': '0.0"%"', 'date': 'yyyy-mm-dd', 'datetime': 'yyyy-mm-dd hh:mm'}


def _excel_value(value, kind):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return timezone.localtime(value).replace(tzinfo=None) if timezone.is_aware(value) else value
    if kind in ('money', 'qty', 'num', 'pct'):
        return float(value)  # openpyxl writes Decimal as text in some versions; floats keep cells numeric
    return value


def to_xlsx(doc):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = doc['title'][:31].replace('/', '-')
    ws.append([AGENCY[1]])
    ws.append([doc['title']])
    ws['A1'].font = Font(bold=True, color=BRAND)
    ws['A2'].font = Font(bold=True, size=14)
    for line in [*_meta_lines(doc), f'Generated {doc["generated_at"]:%Y-%m-%d %H:%M} by {doc["generated_by"]}', *doc['notes']]:
        ws.append([line])
        ws.cell(ws.max_row, 1).font = Font(italic=True, color='555555')
    ws.append([])
    header = ws.max_row + 1
    ws.append([c['label'] for c in doc['columns']])
    for cell in ws[header]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=BRAND)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    rows = [[_excel_value(r.get(c['key']), c['kind']) for c in doc['columns']] for r in doc['rows']]
    total = _total_row(doc)
    formats = [(i, EXCEL_FORMATS[c['kind']]) for i, c in enumerate(doc['columns'], 1) if c['kind'] in EXCEL_FORMATS]
    # Rows are counted here: ws.max_row scans every cell, so reading it per row made big exports quadratic.
    row = header
    for values in rows + ([[_excel_value(v, c['kind']) if i else v for i, (v, c) in enumerate(zip(total, doc['columns']))]]
                          if total else []):
        ws.append(values)
        row += 1
        for i, fmt in formats:
            ws.cell(row, i).number_format = fmt
    if total:
        for i in range(1, len(doc['columns']) + 1):
            ws.cell(row, i).font = Font(bold=True)
    last = header + len(rows)
    ws.freeze_panes = ws.cell(header + 1, 1)
    if rows:
        ws.auto_filter.ref = f'A{header}:{get_column_letter(len(doc["columns"]))}{last}'
    for i, c in enumerate(doc['columns'], 1):
        longest = max([len(c['label'])] + [len(text(r.get(c['key']), c['kind'])) for r in doc['rows'][:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 50)
    out = io.BytesIO()
    wb.save(out)
    return _download(out.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename(doc, 'xlsx'))


WIDE_TEXT = {'details', 'description', 'purpose', 'remarks', 'destination', 'origin', 'vehicle', 'name', 'categories'}
NAME_TEXT = {'ticket_number', 'po_number', 'driver', 'full_name', 'accountable_person', 'supplier', 'service_provider'}


def to_pdf(doc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    brand = colors.HexColor(f'#{BRAND}')
    page = landscape(A4)
    margin = 12 * mm
    cell = ParagraphStyle('cell', fontName='Helvetica', fontSize=7.5, leading=9)
    num = ParagraphStyle('num', parent=cell, alignment=2)
    head = ParagraphStyle('head', parent=cell, fontName='Helvetica-Bold', textColor=colors.white)
    head_num = ParagraphStyle('head_num', parent=head, alignment=2)
    bold, bold_num = ParagraphStyle('b', parent=cell, fontName='Helvetica-Bold'), ParagraphStyle('bn', parent=num, fontName='Helvetica-Bold')
    small = ParagraphStyle('small', fontName='Helvetica', fontSize=8, leading=10, textColor=colors.HexColor('#444444'))

    title_block = Table([[Image(str(LOGO), 16 * mm, 16 * mm), [
        Paragraph(AGENCY[0], ParagraphStyle('a0', fontName='Helvetica', fontSize=8, leading=10)),
        Paragraph(AGENCY[1], ParagraphStyle('a1', fontName='Helvetica-Bold', fontSize=10, leading=12, textColor=brand)),
        Paragraph(escape(doc['title']), ParagraphStyle('t', fontName='Helvetica-Bold', fontSize=14, leading=18)),
    ]]], colWidths=[20 * mm, None])
    title_block.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('LEFTPADDING', (0, 0), (-1, -1), 0)]))
    story = [title_block, Spacer(1, 3 * mm)]
    story += [Paragraph(escape(line), small) for line in _meta_lines(doc)]
    story.append(Spacer(1, 3 * mm))

    cols = doc['columns']
    is_num = [c['kind'] in NUMERIC for c in cols]
    weights = [3.2 if c['key'] in WIDE_TEXT else 2.3 if c['key'] in NAME_TEXT else 1.3 if c['kind'] in NUMERIC
               else 2.2 if c['kind'] == 'datetime' else 1.7 for c in cols]
    width = page[0] - 2 * margin
    col_widths = [width * w / sum(weights) for w in weights]
    p = lambda v, style: Paragraph(escape(v), style)
    data = [[p(c['label'], head_num if n else head) for c, n in zip(cols, is_num)]]
    for r in doc['rows']:
        data.append([p(text(r.get(c['key']), c['kind'], pretty=True), num if n else cell) for c, n in zip(cols, is_num)])
    total = _total_row(doc)
    if total:
        data.append([p(v if i == 0 else text(v, cols[i]['kind'], pretty=True), bold_num if is_num[i] and i else bold)
                     for i, v in enumerate(total)])
    if not doc['rows']:
        data.append([p('No records match these filters.', cell)] + [''] * (len(cols) - 1))
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style = [('BACKGROUND', (0, 0), (-1, 0), brand), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
             ('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.HexColor('#cbd5e1')),
             ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
             ('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3)]
    if total:
        style += [('LINEABOVE', (0, -1), (-1, -1), 1, colors.HexColor('#334155')), ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e8e9fb'))]
    if not doc['rows']:
        style.append(('SPAN', (0, len(data) - 1), (-1, len(data) - 1)))
    table.setStyle(TableStyle(style))
    story.append(table)
    if doc['notes']:
        story.append(Spacer(1, 3 * mm))
        story += [Paragraph(escape(n), small) for n in doc['notes']]

    stamp = f'Generated {doc["generated_at"]:%b %d, %Y %I:%M %p} by {doc["generated_by"]}'

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColor(colors.HexColor('#64748b'))
        canvas.drawString(margin, 7 * mm, f'DOLE Vehicle Records System · {stamp}')
        canvas.drawRightString(page[0] - margin, 7 * mm, f'Page {d.page}')
        canvas.restoreState()

    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=page, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=14 * mm,
                      title=doc['title'], author=AGENCY[1]).build(story, onFirstPage=footer, onLaterPages=footer)
    return _download(out.getvalue(), 'application/pdf', filename(doc, 'pdf'))


EXPORTERS = {'csv': to_csv, 'xlsx': to_xlsx, 'pdf': to_pdf}
