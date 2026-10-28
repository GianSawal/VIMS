"""
Phase 15: DOLE RO3's old workbook (DATA/VEHICLE RECORD.xlsx) -> the standard import workbook.

Old layout:
- sheet RO3: the vehicle master list (rows under the "No. / OFFICE / VEHICLE ..." header) and the authorized
  driver list (rows under "No. / DRIVER");
- one sheet per vehicle, three blocks side by side: maintenance (A-E), fuel and odometer readings (G-I) and the
  RFID / trip log (from column K; its columns are found by header name because their order differs per sheet);
- Sheet20: the fuel purchase order log.

Two steps:
1. prepare(): copy the old workbook with the columns the import needs and the old sheets lack (NEW_* below),
   appended in empty columns only (the logs have balance formulas and merged cells), pre-filled where a guess is
   possible and highlighted where staff must fill them in. The original file is not touched.
2. convert(): read the (prepared and filled) workbook into the standard import workbook. Blank new columns fall back
   to the old behaviour (guess or leave blank), so an unprepared workbook still converts.
Everything inferred or left out is listed on the "Conversion notes" sheet.
The EasyTrip USERNAME / PW columns are never read (the system stores no third-party credentials) and are blanked in
the prepared copy, which gets passed around for filling in.
"""
import re
from datetime import date, datetime, time, timedelta

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from fleet.models import Driver, Vehicle

from .importer import NOTES_SHEET, SHEETS, name_key, plate_key, write_template

# Old spellings (RO3 OFFICE column, driver PLACE column) -> office code. Written to the Offices sheet for review.
OFFICE_ALIASES = {
    'REGIONAL DIRECTOR': 'ORD', 'ORD': 'ORD', 'ARD': 'ARD', 'OARD': 'ARD',
    'REGIONAL OFFICE': 'R03', 'REGIONAL OFFICE III': 'R03', 'RO': 'R03', 'TSSD1': 'R03', 'TSSD2': 'R03', 'MALSU': 'R03',
    'ZAMBALES': 'ZAM', 'ZAM FO': 'ZAM', 'AURORA': 'AUR', 'AU FO': 'AUR', 'BULACAN': 'BUL', 'BUL FO': 'BUL',
    'NUEVA ECIJA': 'NE', 'NE FO': 'NE', 'N.E FO': 'NE', 'BATAAN': 'BAT', 'BAT FO': 'BAT',
    'TARLAC': 'TAR', 'TAR FO': 'TAR', 'PAMP FO': 'PAMP',
}
# Inferred from the model name; anything else is left blank for the reviewer.
VANS = ('HIACE', 'HI ACE', 'L300', 'TRAVIS')
SUVS = ('INNOVA', 'CROSSWIND', 'ADVENTURE', 'REVO', 'SPORTIVO')
DIESEL_ONLY = ('DIESEL', 'DSL', 'HIACE', 'HI ACE', 'L300', 'TRAVIS', 'CROSSWIND')  # models sold here as diesel only
CATEGORY_WORDS = [('PMS', 'Preventive Maintenance'), ('CHANGE OIL', 'Preventive Maintenance'), ('TIRE', 'Tires'),
                  ('BATTERY', 'Battery'), ('AIRCON', 'Air Conditioning'), ('BRAKE', 'Brakes'), ('ENGINE', 'Engine')]
PMS_KM = re.compile(r'PMS\s*\(?\s*([\d,]{3,})\s*(?:KM)?\s*\)?', re.I)  # "PMS (1622)", "PMS 1,287 KM"; not "1000 PMS"
PLATE_LIKE = re.compile(r'\b([A-Z]{3}\s?-?\d{3,4})\b')  # sheet titles carry an earlier plate, e.g. "... PLATE No. SCD2797"

# Columns prepare() adds (header text = what convert() reads).
NEW_MASTER = ['FUEL TYPE', 'VEHICLE TYPE', 'START DATE', 'LTO EXPIRY']  # RO3 vehicle list, after REMARKS
NEW_DRIVER = ['LICENSE TYPE']                                           # RO3 driver list, after MOBILE No.
FUEL_PRICE = 'PRICE/L'                                                  # vehicle sheets, column J (empty, next to FUEL LITERS)
NEW_LOG = ['TIME OUT', 'TIME IN', 'ODO OUT', 'ODO IN']                  # vehicle sheets, after the log's last column
PO_PRICE = 'UNIT PRICE'                                                 # Sheet20 PO log, after DRIVER
TO_FILL = PatternFill('solid', fgColor='FFF2A8')
CREDENTIAL_COLUMNS = ['USERNAME', 'PW']  # RO3 EasyTrip logins: never read, and blanked in the prepared copy


def clean(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v not in (None, '') else ''


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = clean(v).replace('Sept', 'Sep')
    for fmt in ('%d-%b-%Y', '%m/%d/%Y', '%Y-%m-%d', '%B %d, %Y', '%b %d, %Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def to_time(v):
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    if isinstance(v, float) and 0 <= v < 1:  # Excel stores a time as a fraction of a day
        minutes = round(v * 24 * 60)
        return time(minutes // 60, minutes % 60)
    for fmt in ('%H:%M', '%I:%M %p', '%I:%M%p', '%H:%M:%S'):
        try:
            return datetime.strptime(clean(v).upper(), fmt).time()
        except ValueError:
            pass
    return None


def to_num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def has(text, words):
    return any(w in text.upper() for w in words)


def guess_type(desc):
    return 'Van' if has(desc, VANS) else 'SUV' if has(desc, SUVS) else None


def guess_fuel(desc):
    return 'Diesel' if has(desc, DIESEL_ONLY) else None


def header_cols(ws, row):
    """{HEADER TEXT: column number}; the first one wins (SCC1115 has PURPOSE twice)."""
    cols = {}
    for c in ws[row]:
        if c.value not in (None, ''):
            cols.setdefault(clean(c.value).upper(), c.column)
    return cols


def find_row(ws, test, last=None):
    return next((r for r in range(1, (last or ws.max_row) + 1) if test([clean(c.value).upper() for c in ws[r]])), None)


def master_rows(ws):
    """(vehicle header row, driver header row) on RO3."""
    head = find_row(ws, lambda v: v[:2] == ['NO.', 'OFFICE'])
    drivers = find_row(ws, lambda v: v[:2] == ['NO.', 'DRIVER'])
    if not head:
        raise ValueError('Sheet RO3 has no "No. / OFFICE" header row.')
    return head, drivers


def is_vehicle_sheet(ws):
    return ws.title != 'RO3' and any(clean(ws.cell(r, 1).value).startswith('PLATE NO') for r in range(1, 10))


def maint_row(ws):
    return find_row(ws, lambda v: 'PR NO.' in v, 12)


def log_row(ws):
    return find_row(ws, lambda v: 'DATE' in v[10:] and ('CREDIT' in v or 'DEBIT' in v), 12)


def po_row(ws):
    return find_row(ws, lambda v: 'P.O NO.' in v, 12)


# ---------------------------------------------------------------- step 1: add the missing columns

def prepare(source, target):
    """Write a copy of the old workbook with the NEW_* columns; returns notes for the person filling it in."""
    values = load_workbook(source, data_only=True)['RO3']  # for the guesses (VEHICLE is plain text, but be safe)
    wb = load_workbook(source)  # formulas kept
    notes = []

    def add_header(ws, row, col, text):
        cell = ws.cell(row, col, text)
        cell.font, cell.fill = Font(bold=True), TO_FILL
        ws.column_dimensions[cell.column_letter].width = max(ws.column_dimensions[cell.column_letter].width or 0, 14)

    def append_columns(ws, row, names, rows, fill_blank=lambda r, name: True):
        """Add `names` after the last header of `row` (unless already there); highlight blanks on `rows`."""
        cols = header_cols(ws, row)
        start = max(c.column for c in ws[row] if c.value not in (None, '')) + 1  # not max(cols): headers can repeat
        for name in names:
            if name not in cols:
                for r in rows:
                    if ws.cell(r, start).value not in (None, ''):
                        notes.append(f'{ws.title}!{ws.cell(r, start).coordinate} already had "{ws.cell(r, start).value}" '
                                     f'where the new {name} column goes; check it.')
                add_header(ws, row, start, name)
                cols[name] = start
                start += 1
            for r in rows:
                if ws.cell(r, cols[name]).value in (None, '') and fill_blank(r, name):
                    ws.cell(r, cols[name]).fill = TO_FILL
        return cols

    def dropdown(ws, col, rows, choices):
        if rows:
            dv = DataValidation(type='list', formula1='"' + ','.join(label for _, label in choices) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            letter = ws.cell(1, col).column_letter
            dv.add(f'{letter}{min(rows)}:{letter}{max(rows)}')

    ro3 = wb['RO3']
    head, dhead = master_rows(ro3)
    pn = header_cols(ro3, head)['PN']
    vrows = [r for r in range(head + 1, dhead or ro3.max_row + 1) if clean(ro3.cell(r, pn).value)]
    for name in CREDENTIAL_COLUMNS:  # the copy gets passed around: no plaintext EasyTrip logins in it
        col = header_cols(ro3, head).get(name)
        for r in vrows if col else []:
            ro3.cell(r, col).value = None
    desc_col = header_cols(ro3, head).get('VEHICLE')
    cols = append_columns(ro3, head, NEW_MASTER, vrows)
    for r in vrows:
        desc = clean(values.cell(r, desc_col).value) if desc_col else ''
        for name, guess in (('FUEL TYPE', guess_fuel(desc)), ('VEHICLE TYPE', guess_type(desc))):
            cell = ro3.cell(r, cols[name])
            if cell.value in (None, '') and guess:
                cell.value, cell.fill = guess, PatternFill()  # a guess to check, not a blank to fill
        ro3.cell(r, cols['START DATE']).number_format = ro3.cell(r, cols['LTO EXPIRY']).number_format = 'yyyy-mm-dd'
    dropdown(ro3, cols['FUEL TYPE'], vrows, Vehicle.Fuel.choices)
    dropdown(ro3, cols['VEHICLE TYPE'], vrows, Vehicle.Type.choices)
    if dhead:
        name_col = header_cols(ro3, dhead)['DRIVER']
        drows = [r for r in range(dhead + 1, ro3.max_row + 1) if clean(ro3.cell(r, name_col).value)]
        cols = append_columns(ro3, dhead, NEW_DRIVER, drows)
        dropdown(ro3, cols['LICENSE TYPE'], drows, Driver.LicenseType.choices)

    for ws in wb.worksheets:
        if not is_vehicle_sheet(ws):
            continue
        mrow = maint_row(ws)
        if mrow:
            if ws.cell(mrow, 10).value in (None, '', FUEL_PRICE):
                add_header(ws, mrow, 10, FUEL_PRICE)
                for r in range(mrow + 1, ws.max_row + 1):  # a price is needed on every fill-up (rows with liters)
                    if to_num(ws.cell(r, 9).value) and ws.cell(r, 10).value in (None, ''):
                        ws.cell(r, 10).fill = TO_FILL
            else:
                notes.append(f'{ws.title}!J{mrow} is not empty; {FUEL_PRICE} column not added.')
        lrow = log_row(ws)
        if lrow:
            append_columns(ws, lrow, NEW_LOG, range(lrow + 1, ws.max_row + 1), fill_blank=lambda r, name: False)
    po = wb['Sheet20'] if 'Sheet20' in wb.sheetnames else None
    if po and po_row(po):
        append_columns(po, po_row(po), [PO_PRICE], [])
    wb.save(target)
    return notes


# ---------------------------------------------------------------- step 2: convert

class Converter:
    def __init__(self, path):
        self.wb = load_workbook(path, data_only=True)
        self.rows = {name: [] for name in SHEETS}
        self.notes = []
        self.aliases = set()

    def note(self, sheet, row, message):
        self.notes.append((sheet, row, message))

    def office(self, name):
        name = clean(name)
        if name.upper() in OFFICE_ALIASES:
            self.aliases.add(name.upper())
        return name

    # ------------------------------------------------------------ RO3: master list and drivers
    def master(self):
        ws = self.wb['RO3']
        head, drivers_head = master_rows(ws)
        cols = header_cols(ws, head)
        cell = lambda r, name: ws.cell(r, cols[name]).value if name in cols else None
        vehicles = {}
        for r in range(head + 1, drivers_head or ws.max_row + 1):
            plate = clean(cell(r, 'PN')).upper()
            if not plate:
                continue
            vehicles[plate_key(plate)] = {
                'row': r, 'plate': plate, 'office': self.office(cell(r, 'OFFICE')), 'desc': clean(cell(r, 'VEHICLE')).upper(),
                'color': clean(cell(r, 'COLOR')), 'easytrip': clean(cell(r, 'EASYTRIP')).replace(' ', ''),
                'gsis': to_date(cell(r, 'GSIS')), 'lto': to_date(cell(r, 'LTO')), 'renew': clean(cell(r, 'RENEW DATE')),
                'par': clean(cell(r, 'PAR')), 'driver': clean(cell(r, 'DRIVER')), 'status': clean(cell(r, 'REMARKS')),
                'fuel': clean(cell(r, 'FUEL TYPE')), 'type': clean(cell(r, 'VEHICLE TYPE')),
                'start': to_date(cell(r, 'START DATE')), 'lto_expiry': to_date(cell(r, 'LTO EXPIRY')),
            }
        drivers = []
        if drivers_head:
            dcols = header_cols(ws, drivers_head)
            dcell = lambda r, name: ws.cell(r, dcols[name]).value if name in dcols else None
            for r in range(drivers_head + 1, ws.max_row + 1):
                name = clean(dcell(r, 'DRIVER'))
                if name:
                    drivers.append({'row': r, 'name': name, 'license': clean(dcell(r, 'LICENSE NUMBER')),
                                    'expiry': to_date(dcell(r, 'EXP.DATE')), 'office': self.office(dcell(r, 'PLACE')),
                                    'mobile': clean(dcell(r, 'MOBILE NO.')), 'license_type': clean(dcell(r, 'LICENSE TYPE'))})
        return vehicles, drivers

    def match_driver(self, name, drivers):
        """'CALONG, FRANCISCO JR. ZAMORA' or 'NIKKO FERRER' -> the driver list's spelling, when exactly one matches
        on surname + first given name."""
        if ',' in name:
            surname, rest = name.split(',', 1)
            first = (name_key(rest).split() or [''])[0]
        else:
            words = [w for w in name_key(name).split() if w not in ('JR', 'SR', 'III')]
            first, surname = (words[0], words[-1]) if words else ('', '')
        surname = name_key(surname)
        found = [d['name'] for d in drivers if {first, surname} <= set(name_key(d['name']).split())]
        return found[0] if len(found) == 1 else None

    # ------------------------------------------------------------ per-vehicle sheets
    def vehicle_sheets(self, vehicles):
        """{plate key: worksheet}, matched by the sheet's PLATE NO. cell, else by the sheet name."""
        found = {}
        for ws in self.wb.worksheets:
            if not is_vehicle_sheet(ws):
                continue
            plate = next(clean(ws.cell(r, 2).value) for r in range(1, 10) if clean(ws.cell(r, 1).value).startswith('PLATE NO'))
            key = plate_key(plate) if plate_key(plate) in vehicles else plate_key(ws.title)
            if key not in vehicles:
                self.note(ws.title, '', f'Sheet not on the RO3 master list (plate "{plate}"); not converted.')
                continue
            if plate and plate_key(plate) != key:
                self.note(ws.title, '', f'PLATE NO. cell says "{plate}" but the sheet name and RO3 say {vehicles[key]["plate"]}; '
                          f'used {vehicles[key]["plate"]}. Check which is right.')
            found[key] = ws
        return found

    def sheet_driver(self, ws):
        return next((clean(ws.cell(r, 2).value) for r in range(1, 10) if clean(ws.cell(r, 1).value).startswith('DRIVER')), '')

    def maintenance_and_fuel(self, ws, plate):
        """Returns the highest odometer reading seen on the sheet."""
        head = maint_row(ws)
        top = 0
        if not head:
            return top
        price_col = header_cols(ws, head).get(FUEL_PRICE)
        for r in range(head + 1, ws.max_row + 1):
            when, ref, desc, supplier, cost = (ws.cell(r, c).value for c in range(1, 6))
            desc = clean(desc)
            if desc and desc.upper() != 'DESCRIPTION':
                day = to_date(when)
                if not day:
                    self.note(ws.title, r, f'Maintenance "{desc[:40]}" has no readable date; row left out.')
                else:
                    odo = PMS_KM.search(desc)
                    odo = int(odo.group(1).replace(',', '')) if odo else None
                    free = isinstance(cost, str) and cost.strip().upper() == 'FREE'
                    cats = sorted({cat for word, cat in CATEGORY_WORDS if word in desc.upper()})
                    self.rows['Maintenance'].append({
                        'plate_number': plate, 'date': day, 'odometer': odo, 'reference_number': clean(ref),
                        'categories': ', '.join(cats), 'description': desc, 'service_provider': clean(supplier),
                        'parts_cost': to_num(cost) or 0, 'remarks': 'Free of charge (old sheet)' if free else ''})
                    top = max(top, odo or 0)
            day, odo, liters = to_date(ws.cell(r, 7).value), to_num(ws.cell(r, 8).value), to_num(ws.cell(r, 9).value)
            if odo:
                top = max(top, int(odo))
            if day and liters:
                self.rows['Fuel'].append({'plate_number': plate, 'date': day, 'liters': liters,
                                          'odometer': int(odo) if odo else None,
                                          'unit_price': to_num(ws.cell(r, price_col).value) if price_col else None})
        return top

    def toll_log(self, ws, v, drivers):
        head = log_row(ws)
        if not head:
            return
        plate, account = v['plate'], v['easytrip']
        cols = header_cols(ws, head)
        get = lambda r, name: ws.cell(r, cols[name]).value if name in cols else None
        trips = self.trips(ws, v, drivers, head, get)
        prev_bal = opening = last_day = None
        seq, emitted, trip_lines = {}, 0, 0

        def emit(r, day, desc, credit, debit, ticket):
            nonlocal emitted
            if not account:
                self.note(ws.title, r, 'Toll row but the vehicle has no EASYTRIP number on RO3; row left out.')
                return
            n = seq[day] = seq.get(day, -1) + 1  # keeps same-day rows distinct and in sheet order
            self.rows['Toll'].append({
                'provider': 'Easytrip', 'account_number': account, 'plate_number': plate,
                'transacted_at': datetime.combine(day, time(8, 0)) + timedelta(minutes=n), 'description': desc[:200],
                'credit': credit or None, 'debit': debit or None,
                'ticket_number': ticket if ticket in trips else '', 'reference_number': '' if ticket in trips else ticket})
            emitted += 1

        for r in range(head + 1, ws.max_row + 1):
            raw_credit = get(r, 'CREDIT')
            label = clean(raw_credit).upper() if isinstance(raw_credit, str) else ''
            credit, debit, bal = to_num(raw_credit), to_num(get(r, 'DEBIT')), to_num(get(r, 'BALANCE'))
            debit = abs(debit) if debit else debit  # one sheet writes charges as negatives
            day = to_date(get(r, 'DATE'))
            desc, ticket = clean(get(r, 'DESCRIPTION')), clean(get(r, 'TRIP TICKET NO.'))
            reload = None
            if label and 'RELOAD' in label and not debit and bal is None:
                self.note(ws.title, r, f'"{clean(raw_credit)}" written without an amount or balance; row left out.')
                continue
            if label and 'RELOAD' in label:
                if 'TOLL' not in label:  # "RELOAD" with the amount in the DEBIT column
                    credit, debit = debit, None
                else:  # "Toll/Reload 500": a load and a charge on one line
                    m = re.search(r'(\d[\d,]*)', label)
                    reload = float(m.group(1).replace(',', '')) if m else (
                        round(bal - prev_bal + (debit or 0), 2) if bal is not None and prev_bal is not None else None)
            if not credit and not debit and not reload:
                if bal and prev_bal is None and not emitted:
                    opening = bal  # balance carried in from before the log starts
                elif (any(get(r, k) for k in ('PASSENGER', 'PLACE VISIT', 'PURPOSE')) or ticket) and ticket not in trips:
                    trip_lines += 1
                if bal is not None:
                    prev_bal = bal
                continue
            if not day:
                day = last_day
                self.note(ws.title, r, f'No date; used the row above ({day}).' if day else 'No date and none above; row left out.')
                if not day:
                    continue
            last_day = day
            if opening:
                emit(r, day, 'Opening balance carried over from the old spreadsheet', opening, None, '')
                prev_bal, opening = opening, None
            desc = desc or ('RFID load' if credit or reload else 'Toll')
            if reload:
                emit(r, day, f'RFID load ({label.title()})', reload, None, ticket)
            if credit or debit:
                emit(r, day, desc, credit, debit, ticket)
            expected = (prev_bal or 0) + (credit or 0) + (reload or 0) - (debit or 0)
            if bal is not None and prev_bal is not None and abs(expected - bal) > 0.05:  # the sheets drift by a centavo
                self.note(ws.title, r, f'Balance in the sheet is {bal:,.2f} but the rows above give {expected:,.2f} '
                          '(a load or charge is missing in the old log).')
            if bal is not None:
                prev_bal = bal
        if trip_lines:
            self.note(ws.title, '', f'{trip_lines} trip log line(s) without TIME OUT not converted to trips '
                      '(a completed trip needs its times and odometer readings).')

    def trips(self, ws, v, drivers, head, get):
        """Log lines with a TIME OUT become trips (one per ticket); returns the converted ticket numbers."""
        raw_driver = self.sheet_driver(ws)
        driver = raw_driver and (self.match_driver(raw_driver, drivers) or raw_driver)
        done = set()
        for r in range(head + 1, ws.max_row + 1):
            out = to_time(get(r, 'TIME OUT'))
            if not out:
                if get(r, 'TIME OUT') not in (None, ''):
                    self.note(ws.title, r, f'TIME OUT "{get(r, "TIME OUT")}" is not a time (use 08:30); trip left out.')
                continue
            ticket, day = clean(get(r, 'TRIP TICKET NO.')), to_date(get(r, 'DATE'))
            if not ticket or not day:
                self.note(ws.title, r, 'Trip has a TIME OUT but no ticket number or date; left out.')
                continue
            if ticket in done:
                self.note(ws.title, r, f'Ticket {ticket} already converted from a row above; this row\'s times ignored.')
                continue
            back = to_time(get(r, 'TIME IN'))
            departed = datetime.combine(day, out)
            returned = datetime.combine(day, back) if back else None
            if returned and returned < departed:
                returned += timedelta(days=1)  # came back after midnight
            passengers = clean(get(r, 'PASSENGER'))
            self.rows['Trips'].append({
                'ticket_number': ticket, 'plate_number': v['plate'], 'driver_name': driver, 'origin': v['office'],
                'destination': clean(get(r, 'PLACE VISIT')) or clean(get(r, 'DESCRIPTION')), 'purpose': clean(get(r, 'PURPOSE')),
                'departed_at': departed, 'returned_at': returned,
                'odometer_start': int(to_num(get(r, 'ODO OUT'))) if to_num(get(r, 'ODO OUT')) else None,
                'odometer_end': int(to_num(get(r, 'ODO IN'))) if to_num(get(r, 'ODO IN')) else None,
                'remarks': f'Passengers: {passengers}' if passengers else ''})
            done.add(ticket)
        return done

    def po_log(self, vehicles):
        """Sheet20: attach PO number, driver and price to the matching fill-up, or add the fill-up."""
        ws = self.wb['Sheet20'] if 'Sheet20' in self.wb.sheetnames else None
        head = ws and po_row(ws)
        if not head:
            return
        cols = header_cols(ws, head)
        get = lambda r, name: ws.cell(r, cols[name]).value if name in cols else None
        for r in range(head + 1, ws.max_row + 1):
            day, po = to_date(get(r, 'DATE')), clean(get(r, 'P.O NO.'))
            if not day and not po:
                continue
            v = vehicles.get(plate_key(clean(get(r, 'VEHICLE'))))
            if not v or not day:
                self.note(ws.title, r, f'PO log row: vehicle "{clean(get(r, "VEHICLE"))}" not on RO3 or no date; left out.')
                continue
            liters, odo, price = to_num(get(r, 'FUEL LITER')), to_num(get(r, 'ODO')), to_num(get(r, PO_PRICE))
            extra = {'po_number': po, 'driver_name': clean(get(r, 'DRIVER')) or None}
            same = [f for f in self.rows['Fuel'] if f['plate_number'] == v['plate'] and f['date'] == day
                    and (not liters or f['liters'] == liters) and not f.get('po_number')]
            if same:
                same[0].update(extra, unit_price=same[0]['unit_price'] or price)
            else:
                self.rows['Fuel'].append({'plate_number': v['plate'], 'date': day, 'liters': liters,
                                          'odometer': int(odo) if odo else None, 'unit_price': price, **extra})

    # ------------------------------------------------------------ run
    def run(self):
        vehicles, drivers = self.master()
        sheets = self.vehicle_sheets(vehicles)
        for d in drivers:
            self.rows['Drivers'].append({
                'full_name': d['name'], 'office': d['office'], 'contact_number': d['mobile'],
                'license_number': d['license'], 'license_type': d['license_type'] or None, 'license_expiry_date': d['expiry']})
        for key, v in vehicles.items():
            ws, words = sheets.get(key), v['desc'].split()
            top = 0
            remarks = []
            if ws:
                title = ' '.join(clean(c.value).upper() for c in ws[1] if isinstance(c.value, str))
                other = {p for p in PLATE_LIKE.findall(title) if plate_key(p) != key}
                if other:
                    remarks.append(f'Other plate in the old sheet: {", ".join(sorted(other))}')
                top = self.maintenance_and_fuel(ws, v['plate'])
                self.toll_log(ws, v, drivers)
            self.rows['Vehicles'].append({
                'plate_number': v['plate'], 'make': words[0] if words else '', 'model': words[1] if len(words) > 1 else '',
                'variant': ' '.join(words[2:]), 'color': v['color'], 'office': v['office'], 'status': v['status'] or None,
                'vehicle_type': v['type'] or guess_type(v['desc']), 'fuel_type': v['fuel'] or guess_fuel(v['desc']),
                'current_odometer': top or None, 'remarks': '; '.join(remarks)})
            if v['par']:
                driver = v['driver'] and self.match_driver(v['driver'], drivers)
                if v['driver'] and not driver:
                    driver = v['driver']
                    self.note('RO3', v['row'], f'Driver "{v["driver"]}" matches no single name on the driver list; kept as written.')
                self.rows['Assignments'].append({'plate_number': v['plate'], 'office': v['office'], 'start_date': v['start'],
                                                 'accountable_person': v['par'], 'driver_name': driver or None})
            if v['gsis']:
                self.rows['Documents'].append({'plate_number': v['plate'], 'document_type': 'GSIS Insurance', 'expiry_date': v['gsis']})
            if v['lto'] or v['lto_expiry']:
                self.rows['Documents'].append({'plate_number': v['plate'], 'document_type': 'LTO Registration', 'issue_date': v['lto'],
                                               'expiry_date': v['lto_expiry'],
                                               'remarks': f'Renewal month: {v["renew"].title()}' if v['renew'] else ''})
        self.po_log(vehicles)
        self.rows['Offices'] = [{'name_in_old_sheet': a, 'office_code': OFFICE_ALIASES[a]} for a in sorted(self.aliases)]
        self.note('', '', 'Inferred, please check: vehicle type and fuel type where the RO3 columns are blank (HiAce/L300/'
                  'Travis = Van; Innova/Crosswind/Adventure/Revo/Sportivo = SUV; Diesel for models sold only as diesel), '
                  'maintenance categories from the description, the odometer from "PMS (km)" descriptions, trip origin = '
                  'the vehicle\'s office, and the office codes on the Offices sheet.')
        missing = self.missing()
        if missing:
            self.note('', '', 'Still blank (fill them in the prepared old workbook, then convert again): ' + '; '.join(missing) + '.')
        return self

    def missing(self):
        count = lambda sheet, key: sum(1 for r in self.rows[sheet] if r.get(key) in (None, ''))
        checks = [('Fuel', 'unit_price', 'PRICE/L on fill-ups'), ('Assignments', 'start_date', 'START DATE on RO3'),
                  ('Drivers', 'license_type', 'LICENSE TYPE on the driver list'), ('Vehicles', 'fuel_type', 'FUEL TYPE on RO3'),
                  ('Vehicles', 'vehicle_type', 'VEHICLE TYPE on RO3')]
        out = [f'{count(s, k)} {label}' for s, k, label in checks if count(s, k)]
        lto = sum(1 for r in self.rows['Documents'] if r['document_type'] == 'LTO Registration' and not r.get('expiry_date'))
        return out + ([f'{lto} LTO EXPIRY on RO3'] if lto else [])


def convert(source, target):
    """Write the standard import workbook filled from the old workbook; returns {sheet: row count} and the notes."""
    conv = Converter(source).run()
    write_template(target)
    wb = load_workbook(target)
    for sheet, items in conv.rows.items():
        keys = [c['key'] for c in SHEETS[sheet][1]]
        for item in items:
            wb[sheet].append([item.get(k) for k in keys])
    notes = wb.create_sheet(NOTES_SHEET)
    notes.append(['Old sheet', 'Row', 'Note'])
    for n in conv.notes:
        notes.append(list(n))
    notes.column_dimensions['C'].width = 120
    wb.save(target)
    return {sheet: len(items) for sheet, items in conv.rows.items()}, conv.notes
