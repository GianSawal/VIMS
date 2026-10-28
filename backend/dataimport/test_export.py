"""The VEHICLE RECORD export: the form filled from the system, per vehicle sheet, scoped and masked like the API."""
import io

import pytest
from django.contrib.auth.models import Group
from django.core.management import call_command
from openpyxl import load_workbook

from accounts.models import Office, User
from accounts.test_auth import client_for
from audit.models import AuditLog
from dataimport.legacy import convert, master_rows, prepare
from dataimport.test_legacy import fill_new_columns, old_workbook
from fleet.models import Vehicle

pytestmark = pytest.mark.django_db


@pytest.fixture
def records(tmp_path):
    """The synthetic old workbook, imported: one Tarlac vehicle with maintenance, fuel (with PO), toll and a trip."""
    call_command('seed_reference', stdout=io.StringIO())
    tar = Office.objects.create(code='TAR', name='Tarlac Field Office', office_type='FIELD')
    other = Office.objects.create(code='ZAM', name='Zambales Field Office', office_type='FIELD')
    User.objects.create(username='migrator')
    old, ready, out = tmp_path / 'old.xlsx', tmp_path / 'ready.xlsx', tmp_path / 'filled.xlsx'
    old_workbook(old)
    prepare(old, ready)
    fill_new_columns(ready)
    convert(ready, out)
    call_command('import_workbook', str(out), commit=True, user='migrator', out=str(tmp_path / 'runs'), stdout=io.StringIO())
    Vehicle.objects.create(plate_number='ZZZ 999', make='Isuzu', model='Crosswind', vehicle_type='SUV', fuel_type='DIESEL', office=other)
    return tar


def user(name, role, offices=()):
    u = User.objects.create(username=name)
    u.groups.add(Group.objects.get(name=role))
    u.offices.set(offices)
    return u


def download(u, query=''):
    r = client_for(u).get(f'/api/export/vehicle-record/{query}')
    assert r.status_code == 200, r.content
    return load_workbook(io.BytesIO(r.content))


def test_export_fills_the_form_per_vehicle(records):
    wb = download(user('admin', 'System Administrator'))
    assert wb.sheetnames == ['Sheet20', 'RO3', 'SAB 1234', 'ZZZ 999']  # one sheet per vehicle, no empty PLATE NO left
    ro3 = wb['RO3']
    head, dhead = master_rows(ro3)
    row = [c.value for c in ro3[head + 1]]
    assert row[4] == 'SAB 1234' and row[1] == 'Tarlac Field Office' and row[5] == '5400-0000-1111'  # full account number
    assert 'PEDRO S. REYES' in row and 'JUAN P. CRUZ' in row and 'Diesel' in row and 'Van' in row
    assert ro3.cell(head + 1, 5).hyperlink.location == "'SAB 1234'!A1"
    assert [c.value for c in ro3[dhead + 1]][1:7] == ['PEDRO S. REYES', 'C10-00-000001', ro3.cell(dhead + 1, 4).value,
                                                      'Tarlac Field Office', '09170000000', 'Professional']

    ws = wb['SAB 1234']
    assert ws['B4'].value == 'SAB 1234' and ws['E2'].value is not None  # TIRE: last tire job date
    assert [ws.cell(9, c).value for c in (2, 3, 5)] == [None, 'PMS (1,154)', 0]
    assert [ws.cell(9, c).value for c in (8, 9, 10)] == [65, 34.89, 61.5]
    log = [[ws.cell(r, c).value for c in (12, 14, 15)] for r in range(3, 11)]
    assert ['RO3-002', None, None] in log  # the trip line
    assert [None, 1000, None] in log and ['RO3-001', None, 632] in log
    assert ws['P4'].value == '=P3+N4-O4'
    assert [c.value for c in wb['Sheet20'][3]][:7] == [wb['Sheet20']['A3'].value, 'PO-0001', 'SAB 1234', 34.89, 65, 'PEDRO S. REYES', 61.5]
    assert AuditLog.objects.filter(action='export', entity_type='vehicle').exists()


def test_export_is_scoped_and_masked(records):
    clerk = user('clerk', 'Field Office User', [records])
    wb = download(clerk)
    assert wb.sheetnames == ['Sheet20', 'RO3', 'SAB 1234']  # the Zambales vehicle is not theirs
    assert wb['RO3'].cell(master_rows(wb['RO3'])[0] + 1, 6).value == '•••• 1111'  # account number masked without change_rfidaccount
    assert download(user('admin2', 'System Administrator'), f'?office={records.pk}').sheetnames == ['Sheet20', 'RO3', 'SAB 1234']
    nobody = User.objects.create(username='nobody')
    assert client_for(nobody).get('/api/export/vehicle-record/').status_code == 403
