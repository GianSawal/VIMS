"""Phase 15: RO3's old per-vehicle workbook converts into an importable standard workbook (synthetic data)."""
import io
from datetime import datetime, time

import pytest
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from accounts.models import Office, User
from dataimport.importer import SHEETS
from dataimport.legacy import convert, prepare
from fleet.models import Driver, VehicleAssignment, VehicleDocument
from maintenance.models import MaintenanceRecord
from operations.models import FuelRecord, TollTransaction, Trip

pytestmark = pytest.mark.django_db


def old_workbook(path):
    wb = Workbook()
    ro3 = wb.active
    ro3.title = 'RO3'
    ro3.append(['DEPARTMENT OF LABOR AND EMPLOYMENT'])
    ro3.append(['No.', 'OFFICE', 'VEHICLE', 'COLOR', 'PN', 'EASYTRIP', 'USERNAME', 'PW', 'GSIS', 'LTO', 'Renew Date', 'PAR', 'DRIVER', 'REMARKS'])
    ro3.append([1, 'TARLAC', 'TOYOTA HIACE COMMUTER 3.0 M/T', 'WHITE', 'SAB 1234', '5400-0000-1111', 'user1', 'secret-pw',
                datetime(2026, 12, 16), datetime(2026, 1, 19), 'JUNE', 'JUAN P. CRUZ', 'REYES, PEDRO SANTOS', 'SERVICEABLE'])
    ro3.append([])
    ro3.append(['AUTHORIZED DRIVER'])
    ro3.append(['No.', 'DRIVER', 'LICENSE NUMBER', 'EXP.DATE', 'PLACE', 'MOBILE No.'])
    ro3.append([1, 'PEDRO S. REYES', 'C10-00-000001', datetime(2032, 1, 1), 'TAR FO', '09170000000'])

    ws = wb.create_sheet('SAB 1234')
    ws['A1'], ws['H1'] = 'MAINTENANCE RECORD', 'TOYOTA HIACE WHITE WITH PLATE No. SCD 9999'
    for r, (label, value) in enumerate([('VEHICLE :', 'TOYOTA HIACE'), ('COLOR :', 'WHITE'), ('PLATE NO. :', 'SAB1234'),
                                        ('DRIVER :', 'PEDRO REYES'), ('OFFICE :', 'TARLAC FIELD OFFICE')], start=2):
        ws.cell(r, 1, label), ws.cell(r, 2, value)
    for c, h in enumerate(['DATE', 'TRIP TICKET No.', 'DESCRIPTION', 'CREDIT', 'DEBIT', 'BALANCE', 'PASSENGER', 'PLACE VISIT', 'PURPOSE'], start=11):
        ws.cell(2, c, h)
    for c, h in enumerate(['DATE', 'PR No.', 'DESCRIPTION', 'SUPPLIER', 'COST', None, 'DATE', 'ODO', 'FUEL LITERS'], start=1):
        ws.cell(8, c, h)
    ws.append([datetime(2026, 2, 11), None, 'PMS (1,154)', 'Toyota Plaridel', 'Free', None, datetime(2026, 1, 26), 65, 34.89])
    ws.append([datetime(2026, 3, 6), 'IMSD-2026-0049', '2 PCS TIRE', 'CATL TRADING', 14400, None, datetime(2026, 1, 27), 900])
    log = [  # K..P: date, ticket, description, credit, debit, balance
        (None, None, None, None, None, 1000),                                   # opening balance
        (datetime(2026, 3, 2), 'RO3-001', 'Mexico/Tarlac', 'Toll', 632, 368),   # label in CREDIT, charge in DEBIT
        (datetime(2026, 3, 3), 'RO3-002', None, None, None, None),              # trip line only
        ('04-Sept-2026', None, 'SAN FERNANDO NB-X01', None, -96, 272),          # text date, negative charge
        (datetime(2026, 9, 5), None, 'Tarlac/Mexico', 'Toll/Reload/IMSD', 272, 5000),  # reload worked out from the balance
        (datetime(2026, 9, 5), None, 'BAYAD CENTER', 500, None, 5500),
    ]
    for i, row in enumerate(log, start=3):
        for c, v in enumerate(row, start=11):
            ws.cell(i, c, v)
    ws.cell(5, 17, 'JUAN P. CRUZ')  # passenger on the trip-only line
    ws.cell(5, 18, 'Tarlac City')
    ws.cell(5, 19, 'Courtesy call')
    po = wb.create_sheet('Sheet20')
    po.append(['RED PLATE FUEL PURCHASE ORDER LOG'])
    po.append(['DATE', 'P.O No.', 'VEHICLE', 'FUEL LITER', 'ODO', 'DRIVER'])
    po.append([datetime(2026, 1, 26), 'PO-0001', 'SAB 1234', 34.89, 65, 'PEDRO S. REYES'])
    wb.save(path)


def fill_new_columns(path):
    """What staff do in the prepared copy: the columns prepare() added."""
    wb = load_workbook(path)
    ro3, ws, po = wb['RO3'], wb['SAB 1234'], wb['Sheet20']
    ro3['Q3'], ro3['R3'] = datetime(2026, 1, 19), datetime(2027, 6, 30)  # START DATE, LTO EXPIRY
    ro3['G7'] = 'Professional'                                            # LICENSE TYPE
    ws['J9'] = 61.50                                                      # PRICE/L
    ws['T5'], ws['U5'], ws['V5'], ws['W5'] = time(8, 0), '5:30 PM', 900, 1010  # the RO3-002 trip
    po['G3'] = 60
    wb.save(path)


def test_convert_then_import(tmp_path):
    old, out = tmp_path / 'old.xlsx', tmp_path / 'filled.xlsx'
    old_workbook(old)
    counts, notes = convert(old, out)
    assert counts['Vehicles'] == 1 and counts['Drivers'] == 1 and counts['Maintenance'] == 2 and counts['Fuel'] == 1

    wb = load_workbook(out)
    rows = lambda sheet: [dict(zip([c['key'] for c in SHEETS[sheet][1]], r)) for r in wb[sheet].iter_rows(min_row=2, values_only=True)]
    assert not any('secret-pw' in str(v) or 'user1' == v for ws in wb for r in ws.iter_rows(values_only=True) for v in r)

    v = rows('Vehicles')[0]
    assert (v['make'], v['model'], v['vehicle_type'], v['fuel_type'], v['office']) == ('TOYOTA', 'HIACE', 'Van', 'Diesel', 'TARLAC')
    assert v['current_odometer'] == 1154 and 'SCD 9999' in v['remarks']
    assert rows('Assignments')[0]['driver_name'] == 'PEDRO S. REYES'  # "REYES, PEDRO SANTOS" matched to the driver list
    pms, tire = rows('Maintenance')
    assert (pms['odometer'], pms['parts_cost'], pms['categories']) == (1154, 0, 'Preventive Maintenance')
    assert (tire['reference_number'], tire['parts_cost'], tire['categories']) == ('IMSD-2026-0049', 14400, 'Tires')
    assert [(t['credit'], t['debit']) for t in rows('Toll')] == [(1000, None), (None, 632), (None, 96), (5000, None), (None, 272), (500, None)]
    assert len({t['transacted_at'] for t in rows('Toll')}) == 6  # same-day rows stay distinct and ordered
    assert any('1 trip log line' in n[2] for n in notes)
    assert any('1 PRICE/L on fill-ups' in n[2] and '1 START DATE' in n[2] for n in notes)  # what is still blank


def test_prepare_fill_convert_import(tmp_path):
    old, ready, out = tmp_path / 'old.xlsx', tmp_path / 'to-fill.xlsx', tmp_path / 'filled.xlsx'
    old_workbook(old)
    prepare(old, ready)
    wb = load_workbook(ready)
    assert [c.value for c in wb['RO3'][2]][-4:] == ['FUEL TYPE', 'VEHICLE TYPE', 'START DATE', 'LTO EXPIRY']
    assert (wb['RO3']['O3'].value, wb['RO3']['P3'].value) == ('Diesel', 'Van')  # guesses pre-filled for checking
    assert wb['RO3']['G6'].value == 'LICENSE TYPE' and wb['SAB 1234']['J8'].value == 'PRICE/L'
    assert [c.value for c in wb['SAB 1234'][2]][-4:] == ['TIME OUT', 'TIME IN', 'ODO OUT', 'ODO IN']
    assert wb['Sheet20']['G2'].value == 'UNIT PRICE'
    assert load_workbook(old)['RO3'].max_column == 14  # the original is not touched
    assert (wb['RO3']['G3'].value, wb['RO3']['H3'].value) == (None, None)  # EasyTrip login blanked in the copy
    assert load_workbook(old)['RO3']['H3'].value == 'secret-pw'

    fill_new_columns(ready)
    convert(ready, out)
    call_command('seed_reference', stdout=io.StringIO())
    Office.objects.create(code='TAR', name='Tarlac Field Office', office_type='FIELD')
    User.objects.create(username='migrator')
    text = io.StringIO()
    call_command('import_workbook', str(out), commit=True, user='migrator', out=str(tmp_path / 'runs'), stdout=text)
    assert '0 error(s)' in text.getvalue(), text.getvalue()

    assert Driver.objects.get().license_type == Driver.LicenseType.PROFESSIONAL
    assert str(VehicleAssignment.objects.get().start_date) == '2026-01-19'
    assert str(VehicleDocument.objects.get(document_type='LTO_REGISTRATION').expiry_date) == '2027-06-30'
    fuel = FuelRecord.objects.get()
    assert (fuel.unit_price, fuel.po_number, fuel.driver.full_name) == (61.5, 'PO-0001', 'PEDRO S. REYES')
    trip = Trip.objects.get()
    assert (trip.ticket_number, trip.odometer_start, trip.odometer_end, trip.destination) == ('RO3-002', 900, 1010, 'Tarlac City')
    assert trip.driver.full_name == 'PEDRO S. REYES' and (trip.returned_at - trip.departed_at).seconds == 9.5 * 3600
    assert MaintenanceRecord.objects.count() == 2
    t = TollTransaction.objects.order_by('transacted_at')
    assert sum(x.credit - x.debit for x in t) == 5500  # same closing balance as the old sheet
