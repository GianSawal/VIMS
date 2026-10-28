"""Phase 15: the import workbook is validated, dry runs save nothing, re-runs are idempotent, nothing is overwritten."""
import io
import re
from datetime import date

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import load_workbook

from accounts.models import Office, User
from audit.models import AuditLog
from dataimport.importer import SHEETS, plate_key, write_template
from fleet.models import Driver, Vehicle, VehicleAssignment, VehicleDocument
from maintenance.models import MaintenanceRecord
from operations.models import FuelRecord, RFIDAccount, TollTransaction, Trip

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def setup():
    call_command('seed_reference', stdout=io.StringIO())
    Office.objects.create(code='R03', name='DOLE Regional Office III', office_type='REGIONAL')
    Office.objects.create(code='TAR', name='Tarlac Field Office', office_type='FIELD')
    User.objects.create(username='migrator')


def workbook(tmp_path, rows):
    """Fill the real template: rows = {sheet: [{column_key: value}]}."""
    path = tmp_path / 'filled.xlsx'
    write_template(path)
    wb = load_workbook(path)
    for sheet, items in rows.items():
        keys = [c['key'] for c in SHEETS[sheet][1]]
        for item in items:
            wb[sheet].append([item.get(k) for k in keys])
    wb.save(path)
    return path


def run(path, tmp_path, commit=False, user='migrator'):
    out = io.StringIO()
    call_command('import_workbook', str(path), commit=commit, user=user, out=str(tmp_path / 'runs'), stdout=out)
    text = out.getvalue()
    counts = {m[0]: tuple(map(int, m[1:])) for m in re.findall(r'^(\w+)\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)$', text, re.M)}
    report = next((tmp_path / 'runs').glob(f'*-{"commit" if commit else "dryrun"}-report.xlsx'))
    issues = [r for r in load_workbook(report)['Issues'].iter_rows(min_row=2, values_only=True)]
    return counts, issues, text


VEHICLES = [
    {'plate_number': 'sab 1234', 'make': 'Toyota', 'model': 'Innova', 'vehicle_type': 'Van', 'fuel_type': 'diesel',
     'office': 'R03', 'current_odometer': 10000, 'acquisition_cost': '₱1,250,000.00', 'acquisition_date': '03/15/2020'},
    {'plate_number': 'SAC-5678', 'make': 'Mitsubishi', 'model': 'Strada', 'vehicle_type': 'Pick-up', 'fuel_type': 'Diesel',
     'office': 'Tarlac FO', 'status': 'For repair'},
]
GOOD = {
    'Offices': [{'name_in_old_sheet': 'Tarlac FO', 'office_code': 'TAR'}],
    'Vehicles': VEHICLES,
    'Drivers': [{'full_name': 'Emilio C. Capollo', 'office': 'R03', 'license_number': 'n01-23-456789',
                 'license_type': 'Professional', 'license_expiry_date': '2027-05-01'}],
    'Assignments': [
        {'plate_number': 'SAB 1234', 'office': 'R03', 'accountable_person': 'Juan Dela Cruz', 'start_date': '2020-04-01',
         'end_date': '2023-12-31'},
        {'plate_number': 'SAB1234', 'office': 'R03', 'accountable_person': 'Maria Santos', 'driver_name': 'EMILIO C CAPOLLO',
         'start_date': '2024-01-01'},
    ],
    'Documents': [
        {'plate_number': 'SAB 1234', 'document_type': 'LTO Registration', 'reference_number': 'OR-1', 'expiry_date': '2025-03-15'},
        {'plate_number': 'SAB 1234', 'document_type': 'LTO Registration', 'reference_number': 'OR-2', 'expiry_date': '2026-03-15'},
        {'plate_number': 'SAB 1234', 'document_type': 'GSIS Insurance', 'expiry_date': '2026-12-31'},
    ],
    'Maintenance': [{'plate_number': 'SAB 1234', 'date': '2025-06-01', 'odometer': 12000, 'categories': 'Engine, Gearbox',
                     'description': 'Change oil', 'labor_cost': 500, 'parts_cost': '1,500.00'}],
    'Fuel': [
        {'plate_number': 'SAB 1234', 'date': '2025-07-01', 'po_number': 'PO-77', 'liters': 40, 'unit_price': 60,
         'odometer': 12500, 'is_full_tank': 'Yes', 'driver_name': 'Emilio C. Capollo'},
        {'plate_number': 'SAB 1234', 'date': '2025-07-20', 'liters': 30, 'unit_price': 61, 'total_amount': 2000,
         'odometer': 12900, 'driver_name': 'Nobody Known'},
    ],
    'Trips': [{'ticket_number': 'TT-OLD-1', 'plate_number': 'SAB 1234', 'driver_name': 'emilio c. capollo', 'origin': 'San Fernando',
               'destination': 'Tarlac', 'purpose': 'Inspection', 'departed_at': '2025-08-01 08:00', 'returned_at': '2025-08-01 17:00',
               'odometer_start': 12900, 'odometer_end': 13050}],
    'Toll': [{'provider': 'Easytrip', 'account_number': '1234 5678 9012', 'plate_number': 'SAB 1234',
              'transacted_at': '2025-08-01 09:00', 'description': 'NLEX', 'debit': 250, 'ticket_number': 'TT-OLD-1'}],
}


def test_template_has_every_sheet_with_required_marks(tmp_path):
    path = tmp_path / 't.xlsx'
    call_command('import_workbook', template=str(path), stdout=io.StringIO())
    wb = load_workbook(path)
    assert wb.sheetnames == ['Instructions', *SHEETS]
    assert wb['Vehicles']['A1'].value == 'Plate number *'


def test_dry_run_saves_nothing_but_reports(tmp_path):
    counts, issues, text = run(workbook(tmp_path, GOOD), tmp_path)
    assert 'DRY RUN' in text and counts['Vehicles'][0] == 2 and counts['Fuel'][0] == 2
    assert not Vehicle.objects.exists() and not FuelRecord.objects.exists() and not AuditLog.objects.filter(action='import').exists()
    assert list((tmp_path / 'runs').glob('*-source.xlsx'))  # backup copy of the source


def test_commit_imports_every_sheet_and_rerun_creates_nothing(tmp_path):
    path = workbook(tmp_path, GOOD)
    counts, issues, _ = run(path, tmp_path, commit=True)
    assert counts['Vehicles'][0] == 2 and counts['Drivers'][0] == 1 and counts['Assignments'][0] == 2
    v = Vehicle.objects.get(plate_number='SAB 1234')
    assert v.office.code == 'R03' and v.acquisition_cost == 1250000 and v.acquisition_date == date(2020, 3, 15)
    assert Vehicle.objects.get(plate_number='SAC-5678').office.code == 'TAR'  # via the Offices alias sheet
    assert v.current_odometer == 13050  # raised to the highest imported reading
    current = VehicleAssignment.objects.get(vehicle=v, end_date__isnull=True)
    assert current.accountable_person == 'Maria Santos' and current.driver.license_number == 'N01-23-456789'
    lto = VehicleDocument.objects.filter(vehicle=v, document_type='LTO_REGISTRATION')
    assert dict(lto.values_list('reference_number', 'status')) == {'OR-1': 'SUPERSEDED', 'OR-2': 'ACTIVE'}
    m = MaintenanceRecord.objects.get(vehicle=v)
    assert m.status == 'COMPLETED' and m.total_cost == 2000 and sorted(c.name for c in m.categories.all()) == ['Engine', 'Other']
    assert FuelRecord.objects.get(po_number='PO-77').total_amount == 2400
    assert Trip.objects.get(ticket_number='TT-OLD-1').distance_km == 150
    assert TollTransaction.objects.get().trip.ticket_number == 'TT-OLD-1'
    assert RFIDAccount.objects.get().account_number == '123456789012'
    messages = ' '.join(str(i[4]) for i in issues)
    assert 'Unknown category' in messages and 'No driver with this name' in messages and 'Differs from liters x price' in messages
    assert 'raised from 10,000 to 13,050' in messages and '9012' not in ' '.join(str(i[5]) for i in issues if i[3] == 'account_number')[:-4]
    assert AuditLog.objects.get(action='import').changes['summary']['Vehicles']['created'] == 2

    counts, _, _ = run(path, tmp_path, commit=True)
    assert all(c[0] == 0 for c in counts.values()), counts
    assert Vehicle.objects.count() == 2 and FuelRecord.objects.count() == 2 and VehicleDocument.objects.count() == 3


def test_existing_records_are_never_overwritten(tmp_path):
    r03 = Office.objects.get(code='R03')
    v = Vehicle.objects.create(plate_number='SAB 1234', make='Toyota', model='Hiace', vehicle_type='VAN', fuel_type='DIESEL',
                               office=r03, current_odometer=50000, status='UNDER_MAINTENANCE')
    d = Driver.objects.create(full_name='E. Capollo', office=r03, license_number='N01-23-456789', license_type='PROFESSIONAL',
                              license_expiry_date=date(2027, 1, 1), contact_number='0917')
    counts, issues, _ = run(workbook(tmp_path, {'Vehicles': VEHICLES[:1], 'Drivers': GOOD['Drivers']}), tmp_path, commit=True)
    v.refresh_from_db()
    d.refresh_from_db()
    assert (v.model, v.current_odometer, v.status) == ('Hiace', 50000, 'UNDER_MAINTENANCE')  # kept
    assert v.acquisition_cost == 1250000  # blank field filled
    assert counts['Vehicles'][1] == 1 and (d.full_name, d.license_expiry_date) == ('E. Capollo', date(2027, 1, 1))
    assert any(i[0] == 'Vehicles' and 'model' in i[3] and 'kept the values' in i[4] for i in issues)


def test_bad_rows_are_flagged_and_skipped(tmp_path):
    rows = {
        'Vehicles': [VEHICLES[0], {**VEHICLES[0], 'plate_number': 'SAB-1234'},  # same plate, other spelling
                     {**VEHICLES[1], 'plate_number': 'X1', 'acquisition_date': '31/31/2020'},
                     {**VEHICLES[1], 'plate_number': 'X2', 'office': 'Mars Office'},
                     {**VEHICLES[1], 'plate_number': 'X3', 'make': None}],
        'Fuel': [{'plate_number': 'ZZZ 999', 'date': '2025-01-01', 'liters': 1, 'unit_price': 1, 'odometer': 1}],
        'Trips': [{**GOOD['Trips'][0], 'returned_at': None}],
        'Toll': [{**GOOD['Toll'][0], 'credit': 100}],
    }
    counts, issues, _ = run(workbook(tmp_path, rows), tmp_path, commit=True)
    assert counts['Vehicles'][0] == 1 and counts['Vehicles'][3] == 4
    found = {(i[0], i[1]): i[4] for i in issues if i[2] == 'error'}
    assert 'Same plate as row 2' in found[('Vehicles', 3)]
    assert 'Not a date' in found[('Vehicles', 4)]
    assert 'Unknown office' in found[('Vehicles', 5)]
    assert found[('Vehicles', 6)] == 'Required.'
    assert 'No vehicle with this plate' in found[('Fuel', 2)]
    assert 'needs its return time' in found[('Trips', 2)]
    assert 'exactly one of Credit' in found[('Toll', 2)]
    assert Vehicle.objects.count() == 1 and plate_key('sab-1234') == plate_key('SAB 1234')


def test_commit_needs_a_user(tmp_path):
    with pytest.raises(CommandError):
        call_command('import_workbook', str(workbook(tmp_path, GOOD)), commit=True, out=str(tmp_path), stdout=io.StringIO())
